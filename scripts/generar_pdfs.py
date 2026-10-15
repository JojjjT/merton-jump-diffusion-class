"""
Genera los DOCX intermedios de la clase Merton Jump-Diffusion.
Entradas:
- 6 gráficos del laboratorio en assets/laboratorio_graficos
Salidas:
- DOCX en docs/
Nota: la conversión a PDF se realiza con convertir_docx_a_pdf.py.
"""
import os
from docx import Document
from docx.shared import Inches

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LAB_DIR = os.path.join(BASE, "assets", "laboratorio_graficos")
PDF_DIR = os.path.join(BASE, "docs")
os.makedirs(PDF_DIR, exist_ok=True)

def doc_to_docx(doc, docx_path):
    doc.save(docx_path)
    return docx_path

def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")

def create_hoja_trabajo():
    doc = Document()
    doc.add_heading("Hoja de trabajo — ¿GBM o Merton?", 0)
    doc.add_paragraph("Nombre: ________________________________    Fecha: ____________")
    doc.add_paragraph("Objetivo: clasificar trayectorias de precios como GBM (continuo) o Merton (con saltos) y justificar con evidencia.")

    for i in range(1, 7):
        doc.add_heading(f"Gráfico {i}", level=1)
        img = os.path.join(LAB_DIR, f"grafico_{i}.png")
        if os.path.exists(img):
            doc.add_picture(img, width=Inches(6.2))
        doc.add_paragraph("¿GBM o Merton?  ____________________")
        doc.add_paragraph("¿En qué día aproximadamente ocurrió el salto (si lo hubo)?  ____________________")
        doc.add_paragraph("¿Por qué? Evidencia observada:")
        for _ in range(3):
            doc.add_paragraph("________________________________________________________________________")

    doc.add_heading("Preguntas de análisis", level=1)
    doc.add_paragraph("¿Qué señal usé para distinguir GBM de Merton? _______________________________")
    doc.add_paragraph("¿Qué parámetro del modelo (λ, μ_J, σ_J, σ) explicaría el gráfico? ____________")
    doc.add_paragraph("¿Qué error cometí y cómo lo corregí? ______________________________________")

    return doc_to_docx(doc, os.path.join(PDF_DIR, "hoja_trabajo_gbm_o_merton.docx"))

def create_clave_respuestas():
    doc = Document()
    doc.add_heading("Clave de respuestas — ¿GBM o Merton?", 0)
    doc.add_paragraph("Esta clave contiene los parámetros reales usados para generar los 6 gráficos.")
    table = doc.add_table(rows=1, cols=5)
    hdr = table.rows[0].cells
    hdr[0].text = "Gráfico"
    hdr[1].text = "Tipo"
    hdr[2].text = "μ"
    hdr[3].text = "σ"
    hdr[4].text = "Parámetros de salto"
    filas = [
        ("1", "GBM", "0.05", "0.2", "—"),
        ("2", "Merton", "0.05", "0.2", "λ=1.5, μ_J=-0.10, σ_J=0.15"),
        ("3", "GBM", "0.05", "0.2", "—"),
        ("4", "Merton", "0.05", "0.2", "λ=1.5, μ_J=-0.10, σ_J=0.15"),
        ("5", "GBM", "0.05", "0.2", "—"),
        ("6", "Merton", "0.05", "0.2", "λ=1.5, μ_J=-0.10, σ_J=0.15"),
    ]
    for f in filas:
        cells = table.add_row().cells
        for j, v in enumerate(f):
            cells[j].text = v
    doc.add_paragraph("Nota: en los gráficos de Merton el salto es estocástico; no siempre ocurre en el mismo día.")
    return doc_to_docx(doc, os.path.join(PDF_DIR, "clave_respuestas.docx"))

def create_reglas_juego():
    doc = Document()
    doc.add_heading("Instrucciones del juego — El Mercado que Salta", 0)
    doc.add_paragraph("Materiales: 20 cartas, dados y calculadoras.")
    doc.add_heading("Reglas", level=1)
    add_bullets(doc, [
        "Precio base: 100.",
        "Antes de voltear cada carta, predigan el precio final.",
        "Volteen la carta y apliquen el movimiento indicado.",
        "Anoten en qué carta su predicción falló.",
        "Al final, calculen el precio final con la fórmula simplificada.",
    ])
    doc.add_heading("Cartas de difusión (17)", level=1)
    doc.add_paragraph("Tira el dado. Si sale 1–3, el precio sube 2%. Si sale 4–6, baja 2%.")
    doc.add_heading("Cartas de salto (3)", level=1)
    doc.add_paragraph("¡SALTO! El precio cae 15%.")
    doc.add_heading("Fórmula simplificada", level=1)
    doc.add_paragraph("Precio final = Precio anterior × (1 + movimiento de la carta)")
    doc.add_heading("Preguntas de análisis", level=1)
    doc.add_paragraph("¿En qué carta se desvió mi predicción del precio real? ________________________")
    doc.add_paragraph("¿Qué parámetro del modelo (λ o tamaño del salto) representaría ese error? _____")
    return doc_to_docx(doc, os.path.join(PDF_DIR, "reglas_juego_mercado_que_salta.docx"))

def create_taller_salto_real():
    doc = Document()
    doc.add_heading("Taller — Identificar un ‘salto’ en la vida real", 0)
    doc.add_paragraph("Objetivo: aplicar la distinción entre continuidad y discontinuidad a una situación real.")
    doc.add_heading("Instrucciones", level=1)
    add_bullets(doc, [
        "Elijan un evento financiero o económico discontinuo (crisis, anuncio, quiebra, cambio regulatorio).",
        "Describan el contexto y el momento del salto.",
        "Expliquen por qué un modelo de continuidad (GBM) no habría predicho el evento.",
        "Propongan cómo ajustar un modelo con saltos (por ejemplo, aumentar λ o cambiar μ_J).",
        "Señalen qué dato observado los ayudó a identificar el salto.",
    ])
    doc.add_heading("Criterios de entrega", level=1)
    add_bullets(doc, [
        "Claridad del evento y del salto identificado.",
        "Conexión conceptual con el proceso de Merton.",
        "Uso correcto de los parámetros λ, μ_J, σ_J.",
    ])
    return doc_to_docx(doc, os.path.join(PDF_DIR, "tarea_transferencia_salto_real.docx"))

def create_agenda():
    src = os.path.join(BASE, "docs", "00_plan_de_clase.md")
    doc = Document()
    doc.add_heading("Agenda de clase — Saltos por Difusión de Merton", 0)
    with open(src, "r", encoding="utf-8") as f:
        lines = f.read().splitlines()
    for line in lines:
        if line.startswith("# "):
            continue
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.strip() == "" or line.startswith("|"):
            continue
        else:
            doc.add_paragraph(line)
    return doc_to_docx(doc, os.path.join(PDF_DIR, "plan_de_clase_merton.docx"))

def main():
    rutas = [
        create_hoja_trabajo(),
        create_clave_respuestas(),
        create_reglas_juego(),
        create_taller_salto_real(),
        create_agenda(),
    ]
    print("DOCX intermedios generados:")
    for r in rutas:
        print(" -", r)

if __name__ == "__main__":
    main()
